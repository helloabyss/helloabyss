import os
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Turo_Fleet_Tracker.xlsx")

F = "Arial"
BLUE = Font(name=F, color="0000FF")
BLACK = Font(name=F)
GREEN = Font(name=F, color="008000")
BOLD = Font(name=F, bold=True)
HDR = Font(name=F, bold=True, color="FFFFFF")
TITLE = Font(name=F, bold=True, size=14)
NOTE = Font(name=F, italic=True, color="666666", size=9)
HDR_FILL = PatternFill("solid", fgColor="1F2937")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
TOT_FILL = PatternFill("solid", fgColor="E5E7EB")
THIN = Border(bottom=Side(style="thin", color="BBBBBB"))

USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
DATE = "yyyy-mm-dd"
MONTH = "mmm yyyy"
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'

FLEET_N = 75      # car rows
LOG_N = 1000      # trip / expense rows
FR = (5, 5 + FLEET_N - 1)   # fleet data rows
LR = (5, 5 + LOG_N - 1)     # log data rows

wb = Workbook()


def header(ws, row, labels, widths):
    for i, (lab, w) in enumerate(zip(labels, widths), start=1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font, c.fill = HDR, HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = TITLE
    ws["A2"] = sub
    ws["A2"].font = NOTE


def rng(sheet, col, rows):
    return f"{sheet}!${col}${rows[0]}:${col}${rows[1]}"


# ---------------------------------------------------------------- How to use
ws = wb.active
ws.title = "How to use"
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 95
title(ws, "Turo Fleet Tracker — path to $300k/yr net",
      "Educational tool only. Not financial, tax or legal advice. Verify insurance, tax and permits locally.")
rows = [
    ("Step 1", "Assumptions tab: set your goal, start month and targets (yellow cells)."),
    ("Step 2", "Fleet tab: one row per car. Fixed monthly costs (loan, insurance, registration, tracker, parking) go here — NOT in Expenses."),
    ("Step 3", "Trips tab: log every trip from your Turo earnings export. Days, host share, earnings and month fill in automatically."),
    ("Step 4", "Expenses tab: log variable costs (maintenance, tires, repairs, supplies, software). Use Car ID 'FLEET' for overhead."),
    ("Step 5", "Each month: move the Report month on Assumptions forward, then check Car Scorecard and see each car's net, utilization and payback. Cut any car flagged '3 losing months'."),
    ("Step 6", "Dashboard: 12-month P&L, run-rate and how many cars you need at today's net per car."),
    ("Step 7", "Buy Box: test a car before you buy it. Only buy 'BUY' verdicts."),
    ("", ""),
    ("Colour legend", ""),
    ("Blue text, pale yellow fill", "Inputs you type in."),
    ("Bright yellow fill", "Key assumptions that drive the whole workbook."),
    ("Black text", "Formulas — do not overwrite."),
    ("Green text", "Pulled from another tab."),
    ("", ""),
    ("Example rows", "Rows whose Notes say 'EXAMPLE – delete' show the expected format. Clear them before you start."),
    ("Trip timing", "Trips are counted in the month they START. A trip spanning month-end counts fully in the start month."),
    ("Host share", "Trips use the car's earnings-plan share from Fleet; if blank, the default on Assumptions. "
                   "US plans pay 70/80/90% of trip price (Turo help centre, Jan 2026) — re-check, Turo changes terms."),
    ("Not modelled", "Depreciation / resale value, income tax, owner's time. Track resale value separately when you sell a car."),
]
for i, (a, b) in enumerate(rows, start=4):
    ws.cell(row=i, column=1, value=a).font = BOLD
    ws.cell(row=i, column=2, value=b).font = BLACK
    ws.cell(row=i, column=2).alignment = Alignment(wrap_text=True)
ws["A13"].font = Font(name=F, bold=True, color="0000FF"); ws["A13"].fill = INPUT_FILL
ws["A14"].fill = KEY_FILL
ws["A16"].font = Font(name=F, bold=True, color="008000")

# ---------------------------------------------------------------- Assumptions
wa = wb.create_sheet("Assumptions")
title(wa, "Assumptions & targets", "Change the yellow cells. Everything else recalculates.")
wa.column_dimensions["A"].width = 42
wa.column_dimensions["B"].width = 16
wa.column_dimensions["C"].width = 80
A = {}
arows = [
    ("goal", "Annual net profit goal ($)", 300000, USD, "Your target (from the plan)."),
    ("start", "Tracker start month", dt.date(2026, 11, 1), MONTH, "First month on the Dashboard. Use the 1st of a month."),
    ("sel", "Report month (latest closed month)", dt.date(2026, 11, 1), MONTH, "Car Scorecard reports on this month; Dashboard shows months up to it. Move it forward each month. Use the 1st of a month."),
    ("share", "Default host share", 0.80, PCT, "Used when a car's plan is blank. US plans: 70% / 80% / 90% (Turo help centre, earnings plans, Jan 2026)."),
    ("util", "Target utilization", 0.60, PCT, "≈18 booked days a month. Turo's published US average is ~15 days/month (turo.com/list-your-car)."),
    ("net", "Target net per car per month ($)", 350, USD, "Stage-1 gate from the plan. ~$500 is the well-run-fleet benchmark needed for $300k at ~50 cars."),
    ("payback", "Max payback (months)", 30, NUM, "Buy Box rule: car must repay its price from operating cash flow within this many months."),
    ("reserve", "Cash reserve per car (months of costs)", 2, NUM, "Keep this many months of each car's fixed + variable cost in the bank."),
]
for i, (k, lab, val, fmt, note) in enumerate(arows, start=4):
    wa.cell(row=i, column=1, value=lab).font = BOLD
    c = wa.cell(row=i, column=2, value=val)
    c.font, c.fill, c.number_format = BLUE, KEY_FILL, fmt
    wa.cell(row=i, column=3, value=note).font = NOTE
    A[k] = f"Assumptions!$B${i}"

# ---------------------------------------------------------------- Lists (hidden helper)
wl = wb.create_sheet("Lists")
wl["A1"], wl["B1"] = "Expense categories", "Car IDs + FLEET"
cats = ["Maintenance", "Tires", "Repairs / damage", "Cleaning supplies", "Detailing", "Tolls / parking",
        "Fuel / charging", "Software / subscriptions", "Marketing", "Labor / contractor", "Delivery", "Other"]
for i, c in enumerate(cats, start=2):
    wl.cell(row=i, column=1, value=c)
wl["B2"] = "FLEET"
for i in range(FLEET_N):
    wl.cell(row=3 + i, column=2, value=f'=IF(Fleet!A{FR[0]+i}="","",Fleet!A{FR[0]+i})')
wl.sheet_state = "hidden"

# ---------------------------------------------------------------- Fleet
wf = wb.create_sheet("Fleet")
title(wf, "Fleet register", "One row per car. Fixed monthly costs live here. Leave Sold date blank while you own the car.")
fh = ["Car ID", "Year", "Make / model", "Status", "Purchase date", "Sold date", "Purchase price ($)",
      "Cash invested ($)\n(down pmt, or full price if cash)", "Loan payment ($/mo)", "Insurance ($/mo)",
      "Registration ($/yr)", "Tracker & software ($/mo)", "Parking ($/mo)", "Earnings plan host share",
      "Fixed cost ($/mo)", "Notes"]
header(wf, 4, fh, [10, 7, 24, 9, 12, 12, 13, 18, 12, 12, 13, 13, 11, 12, 12, 30])
fleet_ex = [
    ["CAR-01", 2022, "Toyota Corolla LE", "Active", dt.date(2026, 11, 1), None, 18500, 3700, 295, 110, 240, 25, 0, 0.80, None, "EXAMPLE – delete"],
    ["CAR-02", 2021, "Toyota RAV4 Hybrid", "Active", dt.date(2026, 11, 15), None, 24000, 6000, 335, 130, 300, 25, 0, 0.80, None, "EXAMPLE – delete"],
]
fmts = [None, "0", None, None, DATE, DATE, USD, USD, USD, USD, USD, USD, USD, PCT, USD, None]
for r in range(FR[0], FR[1] + 1):
    ex = fleet_ex[r - FR[0]] if r - FR[0] < len(fleet_ex) else [None] * 16
    for col in range(1, 17):
        c = wf.cell(row=r, column=col)
        if col == 15:
            c.value = f'=IF(A{r}="",0,N(I{r})+N(J{r})+N(K{r})/12+N(L{r})+N(M{r}))'
            c.font = BLACK
        else:
            c.value = ex[col - 1]
            c.font, c.fill = BLUE, INPUT_FILL
        if fmts[col - 1]:
            c.number_format = fmts[col - 1]
dv = DataValidation(type="list", formula1='"Active,Sold"', allow_blank=True)
wf.add_data_validation(dv); dv.add(f"D{FR[0]}:D{FR[1]}")
dv = DataValidation(type="list", formula1='"0.7,0.8,0.9"', allow_blank=True)
wf.add_data_validation(dv); dv.add(f"N{FR[0]}:N{FR[1]}")
wf["N4"].comment = Comment("US earnings plans pay 70%, 80% or 90% of trip price (Turo help centre, Jan 2026). "
                           "Some cities use a variable share from 31 Mar 2026 — enter your actual average.", "Tracker")

# ---------------------------------------------------------------- Trips
wt = wb.create_sheet("Trips")
title(wt, "Trip log", "Copy from your Turo earnings export. Trip price = price before Turo's cut. Grey columns are formulas.")
th = ["Trip ID", "Car ID", "Start date", "End date", "Days", "Trip price ($)", "Host share", "Host earnings ($)",
      "Reimbursements & extras ($)\n(tolls, fuel, delivery)", "Turnover cost ($)\n(clean, fuel, labor)",
      "Guest rating (1–5)", "Month", "Net trip ($)", "Notes"]
header(wt, 4, th, [12, 10, 12, 12, 7, 13, 10, 14, 18, 16, 10, 11, 12, 28])
trips_ex = [
    ["T-1001", "CAR-01", dt.date(2026, 11, 3), dt.date(2026, 11, 7), 248, 0, 35, 5],
    ["T-1002", "CAR-01", dt.date(2026, 11, 9), dt.date(2026, 11, 16), 410, 18, 35, 5],
    ["T-1003", "CAR-02", dt.date(2026, 11, 18), dt.date(2026, 11, 22), 312, 0, 35, 5],
    ["T-1004", "CAR-01", dt.date(2026, 11, 20), dt.date(2026, 11, 24), 236, 12, 35, 4],
]
share_lookup = (f'INDEX({rng("Fleet","N",FR)},MATCH(B{{r}},{rng("Fleet","A",FR)},0))')
for r in range(LR[0], LR[1] + 1):
    i = r - LR[0]
    ex = trips_ex[i] if i < len(trips_ex) else None
    vals = {
        1: ex[0] if ex else None, 2: ex[1] if ex else None, 3: ex[2] if ex else None, 4: ex[3] if ex else None,
        6: ex[4] if ex else None, 9: ex[5] if ex else None, 10: ex[6] if ex else None, 11: ex[7] if ex else None,
        14: "EXAMPLE – delete" if ex else None,
    }
    for col in (1, 2, 3, 4, 6, 9, 10, 11, 14):
        c = wt.cell(row=r, column=col, value=vals[col])
        c.font, c.fill = BLUE, INPUT_FILL
    sl = share_lookup.format(r=r)
    wt.cell(row=r, column=5, value=f'=IF(OR(C{r}="",D{r}=""),"",MAX(1,D{r}-C{r}))')
    wt.cell(row=r, column=7, value=f'=IF(B{r}="","",IFERROR(IF(N({sl})>0,{sl},{A["share"]}),{A["share"]}))').font = GREEN
    wt.cell(row=r, column=8, value=f'=IF(OR(F{r}="",G{r}=""),"",F{r}*G{r})')
    wt.cell(row=r, column=12, value=f'=IF(C{r}="","",DATE(YEAR(C{r}),MONTH(C{r}),1))')
    wt.cell(row=r, column=13, value=f'=IF(F{r}="","",N(H{r})+N(I{r})-N(J{r}))')
    for col, fmt in {3: DATE, 4: DATE, 5: NUM, 6: USD, 7: PCT, 8: USD, 9: USD, 10: USD, 11: "0.0", 12: MONTH, 13: USD}.items():
        wt.cell(row=r, column=col).number_format = fmt
    for col in (5, 8, 12, 13):
        wt.cell(row=r, column=col).font = BLACK
dv = DataValidation(type="list", formula1=f"={rng('Fleet','A',FR)}", allow_blank=True)
wt.add_data_validation(dv); dv.add(f"B{LR[0]}:B{LR[1]}")
dv = DataValidation(type="decimal", operator="between", formula1="1", formula2="5", allow_blank=True)
wt.add_data_validation(dv); dv.add(f"K{LR[0]}:K{LR[1]}")

# ---------------------------------------------------------------- Expenses
we = wb.create_sheet("Expenses")
title(we, "Variable expenses", "Do NOT log loan, insurance, registration, tracker or parking here — those are fixed costs on the Fleet tab.")
eh = ["Date", "Car ID\n(or FLEET)", "Category", "Vendor", "Amount ($)", "Notes", "Month"]
header(we, 4, eh, [12, 11, 22, 20, 12, 34, 11])
exp_ex = [
    [dt.date(2026, 11, 2), "FLEET", "Cleaning supplies", "Costco", 64],
    [dt.date(2026, 11, 17), "CAR-01", "Maintenance", "Quick Lube", 79],
    [dt.date(2026, 11, 25), "FLEET", "Software / subscriptions", "Pricing tool", 40],
]
for r in range(LR[0], LR[1] + 1):
    i = r - LR[0]
    ex = exp_ex[i] if i < len(exp_ex) else [None] * 5
    for col in range(1, 7):
        v = ex[col - 1] if col <= 5 else ("EXAMPLE – delete" if ex[0] else None)
        c = we.cell(row=r, column=col, value=v)
        c.font, c.fill = BLUE, INPUT_FILL
    we.cell(row=r, column=7, value=f'=IF(A{r}="","",DATE(YEAR(A{r}),MONTH(A{r}),1))').font = BLACK
    we.cell(row=r, column=1).number_format = DATE
    we.cell(row=r, column=5).number_format = USD2
    we.cell(row=r, column=7).number_format = MONTH
dv = DataValidation(type="list", formula1=f"=Lists!$A$2:$A${1+len(cats)}", allow_blank=True)
we.add_data_validation(dv); dv.add(f"C{LR[0]}:C{LR[1]}")
dv = DataValidation(type="list", formula1=f"=Lists!$B$2:$B${2+FLEET_N}", allow_blank=True)
we.add_data_validation(dv); dv.add(f"B{LR[0]}:B{LR[1]}")

# ---------------------------------------------------------------- Formula helpers
T_CAR, T_MON = rng("Trips", "B", LR), rng("Trips", "L", LR)
E_CAR, E_MON, E_AMT = rng("Expenses", "B", LR), rng("Expenses", "G", LR), rng("Expenses", "E", LR)


def t_sum(col, car, mon):
    return f"SUMIFS({rng('Trips', col, LR)},{T_CAR},{car},{T_MON},{mon})"


def active(fr, mon):
    """1 if the Fleet row fr was owned during month `mon` (first-of-month date)."""
    return (f'IF(AND(Fleet!$A${fr}<>"",N(Fleet!$E${fr})<=EOMONTH({mon},0),'
            f'OR(Fleet!$F${fr}="",N(Fleet!$F${fr})>={mon})),1,0)')


def car_net(fr, car, mon):
    return (f"({t_sum('H', car, mon)}+{t_sum('I', car, mon)}-{t_sum('J', car, mon)}"
            f"-SUMIFS({E_AMT},{E_CAR},{car},{E_MON},{mon})-Fleet!$O${fr}*{active(fr, mon)})")


# ---------------------------------------------------------------- Car Scorecard
wsc = wb.create_sheet("Car Scorecard")
title(wsc, "Car scorecard", "Reports on the Report month set on the Assumptions tab.")
wsc["A3"] = "Month:"; wsc["A3"].font = BOLD
wsc["B3"] = f"={A['sel']}"; wsc["B3"].font = GREEN; wsc["B3"].number_format = MONTH
sh = ["Car ID", "Make / model", "Owned this month", "Booked days", "Utilization", "Earnings + extras ($)",
      "Variable costs ($)", "Fixed cost ($)", "Net this month ($)", "Net last month ($)", "Net 2 months ago ($)",
      "Lifetime net ($)", "Cash invested ($)", "Payback progress", "Flag"]
header(wsc, 5, sh, [10, 22, 10, 10, 11, 14, 13, 12, 13, 13, 14, 13, 13, 12, 24])
M0, M1, M2 = "$B$3", "EDATE($B$3,-1)", "EDATE($B$3,-2)"
for i in range(FLEET_N):
    r, fr = 6 + i, FR[0] + i
    car = f"$A{r}"
    blank = f'$A{r}=""'
    wsc.cell(row=r, column=1, value=f'=IF(Fleet!A{fr}="","",Fleet!A{fr})').font = GREEN
    wsc.cell(row=r, column=2, value=f'=IF(Fleet!A{fr}="","",Fleet!B{fr}&" "&Fleet!C{fr})').font = GREEN
    wsc.cell(row=r, column=3, value=f'=IF({blank},"",{active(fr, M0)})')
    wsc.cell(row=r, column=4, value=f'=IF({blank},"",{t_sum("E", car, M0)})')
    wsc.cell(row=r, column=5, value=f'=IF(OR({blank},N(C{r})=0),"",D{r}/DAY(EOMONTH({M0},0)))')
    wsc.cell(row=r, column=6, value=f'=IF({blank},"",{t_sum("H", car, M0)}+{t_sum("I", car, M0)})')
    wsc.cell(row=r, column=7, value=f'=IF({blank},"",{t_sum("J", car, M0)}+SUMIFS({E_AMT},{E_CAR},{car},{E_MON},{M0}))')
    wsc.cell(row=r, column=8, value=f'=IF({blank},"",Fleet!$O${fr}*C{r})')
    wsc.cell(row=r, column=9, value=f'=IF({blank},"",F{r}-G{r}-H{r})')
    wsc.cell(row=r, column=10, value=f'=IF({blank},"",{car_net(fr, car, M1)})')
    wsc.cell(row=r, column=11, value=f'=IF({blank},"",{car_net(fr, car, M2)})')
    # lifetime: all trips & expenses to date, fixed cost x months owned up to the scorecard month
    end = f'IF(Fleet!$F${fr}="",EOMONTH({M0},0),MIN(Fleet!$F${fr},EOMONTH({M0},0)))'
    months = f'MAX(0,(YEAR({end})-YEAR(Fleet!$E${fr}))*12+MONTH({end})-MONTH(Fleet!$E${fr})+1)'
    wsc.cell(row=r, column=12, value=(
        f'=IF({blank},"",SUMIFS({rng("Trips","M",LR)},{T_CAR},{car},{T_MON},"<="&{M0})'
        f'-SUMIFS({E_AMT},{E_CAR},{car},{E_MON},"<="&{M0})-Fleet!$O${fr}*IF(N(Fleet!$E${fr})=0,0,{months}))'))
    wsc.cell(row=r, column=13, value=f'=IF({blank},"",Fleet!$H${fr})').font = GREEN
    wsc.cell(row=r, column=14, value=f'=IF(OR({blank},N(M{r})=0),"",L{r}/M{r})')
    wsc.cell(row=r, column=15, value=(
        f'=IF(OR({blank},N(C{r})=0),"",IF(AND(I{r}<0,J{r}<0,K{r}<0),"REVIEW: 3 losing months",'
        f'IF(I{r}<{A["net"]},"Below net target",IF(E{r}<{A["util"]},"Low utilization","On track"))))'))
    for col, fmt in {3: "0", 4: NUM, 5: PCT, 6: USD, 7: USD, 8: USD, 9: USD, 10: USD, 11: USD, 12: USD, 13: USD, 14: PCT}.items():
        wsc.cell(row=r, column=col).number_format = fmt
fl = f"O6:O{5+FLEET_N}"
wsc.conditional_formatting.add(fl, FormulaRule(formula=['LEFT(O6,6)="REVIEW"'], font=Font(name=F, bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor="B91C1C")))
wsc.conditional_formatting.add(fl, FormulaRule(formula=['O6="On track"'], font=Font(name=F, color="166534"), fill=PatternFill("solid", fgColor="DCFCE7")))
wsc.conditional_formatting.add(fl, FormulaRule(formula=['OR(O6="Below net target",O6="Low utilization")'], font=Font(name=F, color="92400E"), fill=PatternFill("solid", fgColor="FEF3C7")))

# ---------------------------------------------------------------- Dashboard
wd = wb.create_sheet("Dashboard", 1)
title(wd, "Dashboard — 12-month P&L", "All figures are formulas. Trips count in the month they start.")
FA, FE, FF, FO = (rng("Fleet", c, FR) for c in "AEFO")
dh = ["Month", "Cars owned", "Available car-days", "Booked days", "Utilization", "Trips", "Gross trip price ($)",
      "Earnings + extras ($)", "Variable costs ($)", "Fixed costs ($)", "Net profit ($)", "Net per car ($)",
      "Annualized run-rate ($)", "Cars needed for goal\n(at this net/car)"]
HR = 9
for i, (lab, w) in enumerate(zip(dh, [11, 9, 11, 9, 10, 7, 13, 14, 13, 12, 13, 11, 15, 16]), start=1):
    c = wd.cell(row=HR, column=i, value=lab)
    c.font, c.fill = HDR, HDR_FILL
    c.alignment = Alignment(wrap_text=True, vertical="center")
    wd.column_dimensions[c.column_letter].width = w
wd.row_dimensions[HR].height = 32
D0, D1 = HR + 1, HR + 12
for i in range(12):
    r = D0 + i
    m = f"$A{r}"
    wd.cell(row=r, column=1, value=f"={A['start']}" if i == 0 else f"=EDATE(A{r-1},1)")
    mask = (f'(({FA}<>"")*({FE}<=EOMONTH({m},0))*((({FF}="")+({FF}>={m}))>0))')
    wd.cell(row=r, column=2, value=f"=SUMPRODUCT({mask})")
    wd.cell(row=r, column=3, value=f"=B{r}*DAY(EOMONTH({m},0))")
    wd.cell(row=r, column=4, value=f"=SUMIFS({rng('Trips','E',LR)},{T_MON},{m})")
    wd.cell(row=r, column=5, value=f"=IF(C{r}=0,0,D{r}/C{r})")
    wd.cell(row=r, column=6, value=f"=COUNTIFS({T_MON},{m})")
    wd.cell(row=r, column=7, value=f"=SUMIFS({rng('Trips','F',LR)},{T_MON},{m})")
    wd.cell(row=r, column=8, value=f"=SUMIFS({rng('Trips','H',LR)},{T_MON},{m})+SUMIFS({rng('Trips','I',LR)},{T_MON},{m})")
    wd.cell(row=r, column=9, value=f"=SUMIFS({rng('Trips','J',LR)},{T_MON},{m})+SUMIFS({E_AMT},{E_MON},{m})")
    wd.cell(row=r, column=10, value=f"=SUMPRODUCT({mask}*{FO})")
    wd.cell(row=r, column=11, value=f"=H{r}-I{r}-J{r}")
    wd.cell(row=r, column=12, value=f"=IF(B{r}=0,0,K{r}/B{r})")
    wd.cell(row=r, column=13, value=f"=K{r}*12")
    wd.cell(row=r, column=14, value=f'=IF(L{r}<=0,"n/a",{A["goal"]}/12/L{r})')
    # months after the report month stay blank so totals cover closed months only
    for col in range(2, 15):
        c = wd.cell(row=r, column=col)
        c.value = f'=IF({m}>{A["sel"]},"",{c.value[1:]})'
    for col, fmt in {1: MONTH, 2: NUM, 3: NUM, 4: NUM, 5: PCT, 6: NUM, 7: USD, 8: USD, 9: USD, 10: USD,
                     11: USD, 12: USD, 13: USD, 14: NUM1}.items():
        wd.cell(row=r, column=col).number_format = fmt
        wd.cell(row=r, column=col).font = GREEN if (col == 1 and i == 0) else BLACK
        wd.cell(row=r, column=col).border = THIN
TR = D1 + 1
wd.cell(row=TR, column=1, value="12-mo total").font = BOLD
for col in (3, 4, 6, 7, 8, 9, 10, 11):
    L = wd.cell(row=D0, column=col).column_letter
    wd.cell(row=TR, column=col, value=f"=SUM({L}{D0}:{L}{D1})")
wd.cell(row=TR, column=5, value=f"=IF(C{TR}=0,0,D{TR}/C{TR})")
for col in range(1, 15):
    c = wd.cell(row=TR, column=col)
    c.fill = TOT_FILL
    c.font = BOLD
    c.number_format = wd.cell(row=D0, column=col).number_format
wd.conditional_formatting.add(f"E{D0}:E{D1}", CellIsRule(operator="lessThan", formula=[A["util"]], font=Font(name=F, color="B91C1C")))
wd.conditional_formatting.add(f"K{D0}:K{D1}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, color="B91C1C")))

# KPI tiles
tiles = [
    ("A", "Annual goal", f"={A['goal']}", USD),
    ("C", "Net profit (12 mo)", f"=K{TR}", USD),
    ("E", "% of goal", f"=IF({A['goal']}=0,0,K{TR}/{A['goal']})", PCT),
    ("G", "Report-month run-rate", f"=IFERROR(INDEX(M{D0}:M{D1},MATCH({A['sel']},A{D0}:A{D1},0)),0)", USD),
    ("I", "Cars owned (that month)", f"=IFERROR(INDEX(B{D0}:B{D1},MATCH({A['sel']},A{D0}:A{D1},0)),0)", NUM),
    ("K", "Cars needed at that net/car", f'=IFERROR(INDEX(N{D0}:N{D1},MATCH({A["sel"]},A{D0}:A{D1},0)),"n/a")', NUM1),
    ("M", "Cash reserve to hold", f"=IFERROR(INDEX(B{D0}:B{D1},MATCH({A['sel']},A{D0}:A{D1},0))*{A['reserve']}*"
                                   f"(INDEX(I{D0}:I{D1},MATCH({A['sel']},A{D0}:A{D1},0))+INDEX(J{D0}:J{D1},MATCH({A['sel']},A{D0}:A{D1},0)))"
                                   f"/MAX(1,INDEX(B{D0}:B{D1},MATCH({A['sel']},A{D0}:A{D1},0))),0)", USD),
]
for col, lab, f, fmt in tiles:
    wd[f"{col}4"] = lab
    wd[f"{col}4"].font = NOTE
    wd[f"{col}5"] = f
    wd[f"{col}5"].font = Font(name=F, bold=True, size=16)
    wd[f"{col}5"].number_format = fmt
wd["A7"] = "Rows show months up to the Report month (Assumptions tab). Tiles use the Report month. Cars bought mid-month carry a full month of fixed cost."
wd["A7"].font = NOTE

ch = BarChart()
ch.type = "col"
ch.title = "Net profit by month"
ch.y_axis.title = "$"
ch.y_axis.numFmt = "$#,##0"
ch.add_data(Reference(wd, min_col=11, min_row=HR, max_row=D1), titles_from_data=True)
ch.set_categories(Reference(wd, min_col=1, min_row=D0, max_row=D1))
ch.legend = None
ch.height, ch.width = 7.5, 22
wd.add_chart(ch, f"A{TR + 2}")

# ---------------------------------------------------------------- Buy Box
wbx = wb.create_sheet("Buy Box")
title(wbx, "Buy Box — test a car before you buy it",
      "Fill a row per candidate with numbers from comparable Turo listings in your city. BUY = pays back fast enough AND clears the net target after the loan.")
bh = ["Candidate", "Price ($)", "Cash invested ($)", "Loan payment ($/mo)", "Expected daily price ($)",
      "Expected booked days / mo", "Host share", "Avg trip length (days)", "Turnover cost ($/trip)",
      "Maintenance ($/mo)", "Insurance + reg + tracker ($/mo)", "Host earnings ($/mo)", "Turnover costs ($/mo)",
      "Operating cash flow ($/mo, before loan)", "Net after loan ($/mo)", "Payback (months)", "Verdict"]
header(wbx, 4, bh, [24, 10, 11, 11, 11, 11, 9, 10, 11, 11, 13, 12, 12, 14, 12, 10, 9])
buy_ex = [
    ["2022 Toyota Corolla (EXAMPLE)", 18000, 3600, 290, 62, 19, 0.80, 3.5, 30, 75, 120],
    ["2021 Toyota RAV4 Hybrid (EXAMPLE)", 24000, 6000, 330, 78, 20, 0.80, 4, 30, 80, 130],
]
for i in range(15):
    r = 5 + i
    ex = buy_ex[i] if i < len(buy_ex) else [None] * 11
    for col in range(1, 12):
        c = wbx.cell(row=r, column=col, value=ex[col - 1])
        c.font, c.fill = BLUE, INPUT_FILL
    b = f'$A{r}=""'
    wbx.cell(row=r, column=12, value=f'=IF({b},"",E{r}*F{r}*G{r})')
    wbx.cell(row=r, column=13, value=f'=IF({b},"",IF(N(H{r})=0,0,F{r}/H{r}*I{r}))')
    wbx.cell(row=r, column=14, value=f'=IF({b},"",L{r}-M{r}-J{r}-K{r})')
    wbx.cell(row=r, column=15, value=f'=IF({b},"",N{r}-D{r})')
    wbx.cell(row=r, column=16, value=f'=IF({b},"",IF(N{r}<=0,"never",B{r}/N{r}))')
    wbx.cell(row=r, column=17, value=(
        f'=IF({b},"",IF(AND(ISNUMBER(P{r}),P{r}<={A["payback"]},O{r}>={A["net"]}),"BUY","PASS"))'))
    for col, fmt in {2: USD, 3: USD, 4: USD, 5: USD, 6: NUM, 7: PCT, 8: "0.0", 9: USD, 10: USD, 11: USD,
                     12: USD, 13: USD, 14: USD, 15: USD, 16: NUM1}.items():
        wbx.cell(row=r, column=col).number_format = fmt
vr = "Q5:Q19"
wbx.conditional_formatting.add(vr, FormulaRule(formula=['Q5="BUY"'], font=Font(name=F, bold=True, color="166534"), fill=PatternFill("solid", fgColor="DCFCE7")))
wbx.conditional_formatting.add(vr, FormulaRule(formula=['Q5="PASS"'], font=Font(name=F, bold=True, color="B91C1C"), fill=PatternFill("solid", fgColor="FEE2E2")))
wbx["A21"] = ("Payback = price ÷ operating cash flow. Ignores resale value (upside) and depreciation (downside). "
              "Rules come from the Assumptions tab.")
wbx["A21"].font = NOTE

# fonts everywhere default to Arial
for s in wb.worksheets:
    for row in s.iter_rows():
        for c in row:
            if c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)

wb.save(OUT)
print("saved", OUT)
